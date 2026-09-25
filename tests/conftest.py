import re
import uuid
import pytest
from fastapi.testclient import TestClient
from app.config import Settings
from app.main import create_app
from app.db import now, uid
from app.security import hash_password


@pytest.fixture
def workspace(tmp_path):
    name = "skillsprint_test_" + uuid.uuid4().hex
    settings = Settings(
        _env_file=None,
        app_secret_key="t" * 48,
        mongodb_uri="mongodb://localhost:27017",
        mongodb_db_name=name,
        upload_dir=tmp_path / "uploads",
        openai_api_key="",
        run_worker=False,
    )
    app = create_app(settings)
    with TestClient(app) as client:
        db = app.state.db
        password = "Test-password-123!"
        for role in ["admin", "training_manager", "reviewer", "manager", "employee"]:
            db.users.insert_one(
                {
                    "_id": role,
                    "email": role + "@test.local",
                    "name": role.title(),
                    "password_hash": hash_password(password),
                    "role": role,
                    "active": True,
                    "created_at": now(),
                }
            )
        yield client, db, settings
        assert db.name.startswith("skillsprint_test_") and len(db.name) > 20
        db.client.drop_database(db.name)


def csrf(html):
    return re.search(r'name="csrf_token" value="([^"]+)"', html).group(1)


def login(client, role="admin"):
    page = client.get("/login")
    if 'name="email"' not in page.text:
        client.post("/logout", data={"csrf_token": csrf(page.text)})
        page = client.get("/login")
    response = client.post(
        "/login",
        data={
            "email": role + "@test.local",
            "password": "Test-password-123!",
            "csrf_token": csrf(page.text),
        },
    )
    assert response.status_code == 200
    assert "A better first chapter" in response.text
    return csrf(response.text)


def upload_docx(
    client,
    token,
    text="Employees must lock their workstations before leaving them.",
    document_id="TEST-01",
    version="1.0",
    effective="2026-01-01",
):
    import io
    from docx import Document

    doc = Document()
    doc.add_heading("Security procedure", 0)
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "All employees"
    stream = io.BytesIO()
    doc.save(stream)
    return client.post(
        "/documents/upload",
        data={
            "csrf_token": token,
            "document_id": document_id,
            "title": "Security procedure",
            "version": version,
            "effective_date": effective,
            "category": "Policy",
            "role_id": "",
        },
        files={
            "file": (
                "policy.docx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
