import json, io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock
import httpx, pytest
from openai import APITimeoutError, RateLimitError, AuthenticationError
from app.generation import structured, GenerationFailure
from app.schemas import FullOnboardingPlan
from app.ingestion import parse
from .conftest import login

PACK = Path(__file__).resolve().parents[1] / "sample_documents" / "asterbridge"
CASES = json.loads((PACK / "reference/adversarial_cases.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("case", CASES, ids=lambda c: c["case_id"])
def test_all_adversarial_pdf_fixtures_are_quarantined(workspace, case):
    client, db, _ = workspace
    token = login(client)
    users = list(db.users.find())
    path = PACK / case["source_path"]
    response = client.post(
        "/documents/upload",
        data={
            "csrf_token": token,
            "document_id": case["case_id"],
            "title": case["scenario"],
            "version": "1",
            "effective_date": "2026-01-01",
        },
        files={"file": (path.name, path.read_bytes(), "application/pdf")},
    )
    assert response.status_code == 200
    doc = db.documents.find_one({"document_id": case["case_id"]})
    assert doc["suspicious"]
    assert (
        client.post(
            "/documents/" + doc["_id"] + "/activate", data={"csrf_token": token}
        ).status_code
        == 422
    )
    assert (
        client.post("/documents/" + doc["_id"] + "/extract", data={"csrf_token": token}).status_code
        == 422
    )
    assert list(db.users.find()) == users
    assert db.plans.count_documents({}) == 0 and db.jobs.count_documents({}) == 0
    assert db.learning_progress.count_documents({}) == 0


@pytest.mark.parametrize(
    "outcome,expected_calls",
    [("timeout", 2), ("quota", 1), ("auth", 1), ("invalid", 2), ("refused", 1)],
)
def test_provider_failures_are_bounded_and_do_not_expose_secrets(
    workspace, monkeypatch, outcome, expected_calls
):
    _, _, settings = workspace
    settings = settings.model_copy(update={"openai_api_key": "SYNTHETIC_SECRET_DO_NOT_EXPOSE"})
    request = httpx.Request("POST", "https://example.invalid")
    response = httpx.Response(429 if outcome == "quota" else 401, request=request)
    failure = {
        "timeout": APITimeoutError(request=request),
        "quota": RateLimitError("secret", response=response, body=None),
        "auth": AuthenticationError("secret", response=response, body=None),
        "invalid": ValueError("secret"),
    }
    provider = Mock()
    provider.responses.parse.side_effect = failure.get(outcome)
    if outcome == "refused":
        provider.responses.parse.return_value = SimpleNamespace(output_parsed=None)
    monkeypatch.setattr("app.generation.OpenAI", lambda **kwargs: provider)
    monkeypatch.setattr("app.generation.time.sleep", lambda _: None)
    with pytest.raises(GenerationFailure) as error:
        structured(
            settings, FullOnboardingPlan, "generate_v3.txt", {"employee": {}, "requirements": []}
        )
    assert "SYNTHETIC_SECRET" not in str(error.value) and "secret" not in str(error.value)
    assert provider.responses.parse.call_count == expected_calls
    provider.close.assert_called_once()


def test_unfamiliar_docx_tables_and_long_paragraphs_preserve_all_text():
    from docx import Document

    doc = Document()
    text = "Analysts must record a fictional observatory calibration. " * 140
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = (
        "If the lens is unavailable, technicians must record a deferred inspection."
    )
    stream = io.BytesIO()
    doc.save(stream)
    blocks = parse(stream.getvalue(), "unfamiliar.docx")
    assert "".join(b["text"] for b in blocks if b["section_id"].startswith("para")) == text
    assert len({b["section_id"] for b in blocks}) == len(blocks)
    assert any("deferred inspection" in b["text"] for b in blocks)


def test_docx_expansion_limit_and_macro_rejection():
    import zipfile

    for name, payload in [
        ("word/vbaProject.bin", b"x"),
        ("word/document.xml", b"x" * (25 * 1024 * 1024 + 1)),
    ]:
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(name, payload)
        with pytest.raises(ValueError):
            parse(stream.getvalue(), "unsafe.docx")


CONFLICTS = json.loads((PACK / "reference/conflict_cases.json").read_text(encoding="utf-8"))
VERSIONS = json.loads((PACK / "reference/version_changes.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("case", CONFLICTS, ids=lambda c: c["case_id"])
def test_ten_reviewed_conflict_values_are_detected(case):
    from app.policy import conflict_groups

    blocks = parse((PACK / case["source_path"]).read_bytes(), "conflict.pdf")
    assert blocks and not any(b["suspicious"] for b in blocks)
    # Explicit reviewed fixture annotations; this does not claim automatic semantic extraction.
    requirements = []
    for index, value in enumerate([case["authoritative_value"], case["conflicting_value"]]):
        requirements.append(
            {
                "_id": str(index),
                "requirement_id": str(index),
                "text": blocks[0]["text"],
                "policy_rule": {
                    "key": case["disputed_field"],
                    "value": str(value),
                    "condition": "",
                    "exception": "",
                },
            }
        )
    groups = conflict_groups(requirements, "test-role")
    assert len(groups) == 1 and groups[0]["key"] == case["disputed_field"]


@pytest.mark.parametrize("case", VERSIONS, ids=lambda c: c["change_id"])
def test_ten_actual_policy_versions_invalidate_linked_plans(workspace, case):
    client, db, _ = workspace
    token = login(client)
    for version, path, date in [
        (case["old_version"], PACK / case["historical_source"], "2026-01-01"),
        (case["new_version"], PACK / "current" / (case["document_id"] + "_v2_0.pdf"), "2026-09-01"),
    ]:
        result = client.post(
            "/documents/upload",
            data={
                "csrf_token": token,
                "document_id": case["document_id"],
                "title": "Version test",
                "version": version,
                "effective_date": date,
            },
            files={"file": (path.name, path.read_bytes(), "application/pdf")},
        )
        assert result.status_code == 200
        doc = db.documents.find_one({"document_id": case["document_id"], "version": version})
        assert (
            client.post(
                "/documents/" + doc["_id"] + "/activate", data={"csrf_token": token}
            ).status_code
            == 200
        )
        if version == case["old_version"]:
            db.plans.insert_one(
                {
                    "_id": "old",
                    "employee_id": "e",
                    "status": "Published",
                    "source_document_ids": [doc["_id"]],
                }
            )
    assert db.plans.find_one({"_id": "old"})["status"] == "Stale sources"
    assert (
        db.documents.find_one({"document_id": case["document_id"], "version": case["old_version"]})[
            "status"
        ]
        == "superseded"
    )
