"""Fresh unfamiliar DOCX/role rehearsal and three real four-requirement timings.
This authored rehearsal is not the competition organizer's unseen test pack.
"""
import sys,io,json,time
from pathlib import Path
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from app.config import ROOT,Settings
from app.db import connect,initialize
from app.ingestion import ingest
from app.worker import enqueue,process_one

def main():
    settings=Settings().model_copy(update={'mongodb_db_name':'skillsprint_phase4_hidden_'+uuid4().hex,'run_worker':False,'upload_dir':ROOT/'tmp'/'hidden_ready'})
    client,db=connect(settings);initialize(db)
    out=ROOT/'hidden_test_ready';out.mkdir(exist_ok=True)
    doc=Document();doc.add_heading('Fictional Observatory Calibration Procedure',0)
    for text in ['Calibration coordinators must record the lens serial number before each simulated alignment.',
                 'When a lens is unavailable, coordinators must record a deferred inspection before closing the shift.',
                 'Coordinators must use fictional sensor readings during onboarding practice.']:
        doc.add_paragraph(text)
    table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='Supervisors must review the alignment log before a coordinator works independently.'
    file=out/'observatory_rehearsal.docx';doc.save(file)
    report={'scope':'Authored unfamiliar role/DOCX rehearsal; automated fixture approval, not human verification. Three four-requirement real generations; other local evaluation work may run concurrently.','runs':[]}
    try:
        source=ingest(db,settings,file.read_bytes(),file.name,{'document_id':'OBS-01','version':'1','title':'Observatory calibration','effective_date':'2026-01-01','category':'SOP','role_ids':['observatory']},'evaluation')
        db.documents.update_one({'_id':source['_id']},{'$set':{'status':'active'}})
        db.requirements.update_many({'document_id':source['_id']},{'$set':{'status':'approved'}})
        assert db.requirements.count_documents({})==4
        db.employees.insert_one({'_id':'learner','name':'Fictional rehearsal learner','role_id':'observatory','department':'Calibration','experience':'Beginner','joining_date':'2026-09-25'})
        for number in range(3):
            job=enqueue(db,'evaluation','generate','learner');start=time.perf_counter();process_one(db,settings)
            record=db.jobs.find_one({'_id':job['_id']});plan=db.plans.find_one({'_id':job['_id']})
            report['runs'].append({'run':number+1,'seconds':round(time.perf_counter()-start,3),'job':record,'plan':plan})
            print({'run':number+1,'status':record['status'],'core_passed':plan['validation']['core_passed'] if plan else False},flush=True)
    finally:
        (ROOT/'reports/phase4/hidden_ready.json').write_text(json.dumps(report,indent=2,default=str),encoding='utf-8')
        assert db.name.startswith('skillsprint_phase4_hidden_');client.drop_database(db.name);client.close()
if __name__=='__main__':main()
