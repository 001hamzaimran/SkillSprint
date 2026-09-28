"""Isolated, fictional React video workspace; saved AI output is explicitly a replay."""
import json
import sys
from copy import deepcopy
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import uvicorn
from docx import Document
from app.config import ROOT, Settings
from app.db import connect, initialize, now, audit
from app.main import create_app
from app.security import hash_password
from app.policy import fingerprint
from app.learning import learning_checks
from app.schemas import FullOnboardingPlan
from app.validation import validate_plan


def main():
    base = Settings()
    parsed = urlsplit(base.database_url)
    if parsed.hostname not in ('localhost', '127.0.0.1'):
        raise SystemExit('This recording helper only accepts a local PostgreSQL server.')
    schema = 'skillsprint_video_' + uuid4().hex
    settings = base.model_copy(update={
        'database_url': urlunsplit(parsed._replace(netloc=parsed.netloc.replace(':5432', ':5433'), path='/postgres')),
        'postgres_schema': schema, 'mongodb_db_name': schema,
        'app_env': 'development', 'session_cookie_secure': False,
        'app_base_url': 'http://127.0.0.1:8001', 'frontend_url': 'http://127.0.0.1:8001/app',
        'run_worker': False, 'openai_api_key': '', 'smtp_host': '',
        'upload_dir': ROOT / 'tmp' / schema,
    })
    client, db = connect(settings)
    assert db.name == schema and db.name.startswith('skillsprint_video_')
    try:
        initialize(db)
        settings.upload_dir.mkdir(parents=True, exist_ok=True)
        for role, name in [('admin', 'Demo Administrator'), ('employee', 'Ayesha Khan (Demo)'), ('reviewer', 'Demo Reviewer')]:
            db.users.insert_one({'_id': role, 'email': role + '@demo.local', 'name': name, 'role': role,
                                 'active': True, 'password_hash': hash_password('Local-Demo-2026!'), 'created_at': now()})
        db.job_roles.insert_one({'_id': 'support', 'name': 'Customer Support Specialist', 'department': 'Customer Operations', 'description': 'Fictional customer service and incident handling role.'})
        db.job_roles.insert_one({'_id': 'operations', 'name': 'Operations Coordinator', 'department': 'Operations', 'description': 'Fictional logistics coordination role.'})
        employee = {'_id': 'learner', 'name': 'Ayesha Khan (Demo)', 'role_id': 'support',
                    'role_name': 'Customer Support Specialist', 'department': 'Customer Operations',
                    'experience': 'Beginner', 'joining_date': '2026-09-28', 'manager_id': 'admin',
                    'user_id': 'employee', 'created_at': now()}
        db.employees.insert_one(employee)
        evidence = json.loads((ROOT / 'reports/phase3_live_smoke.json').read_text(encoding='utf-8'))
        content = deepcopy(evidence['initial']['content'])
        content['title'] = 'Customer Support — Evidence-Based Onboarding'
        content['summary'] = 'LOCAL DEMO: saved real AI output, replayed with fictional accounts. No fresh AI request is made.'
        content['role_id'] = 'support'
        requirements = []
        for index, item in enumerate(content['items']):
            item['role_id'] = 'support'
            did, sid, rid = item['source_document_id'], item['source_section_id'], item['requirement_id']
            title = 'Workstation Security Policy' if 'lock' in item['source_quote'].lower() else 'Customer Incident Response Policy'
            path = settings.upload_dir / (did + '.docx')
            doc = Document(); doc.add_heading(title, 0); doc.add_paragraph(item['source_quote']); doc.save(path)
            db.documents.insert_one({'_id': did, 'document_id': f'DEMO-POL-{index+1:02d}', 'version': '1.0',
                'title': title, 'status': 'active', 'digest': did, 'filename': path.name, 'path': str(path),
                'effective_date': '2026-09-01', 'category': 'Policy', 'role_ids': ['support'], 'suspicious': False, 'created_at': now()})
            db.source_sections.insert_one({'_id': did + sid, 'document_id': did, 'section_id': sid,
                'text': item['source_quote'], 'heading': title, 'location': 'Saved AI evidence, fictional policy', 'suspicious': False})
            requirement = {'_id': rid, 'requirement_id': rid, 'document_id': did, 'section_id': sid,
                'text': item['source_quote'], 'title': title, 'mandatory': item['mandatory'], 'status': 'approved',
                'role_ids': ['support'], 'due_stage': 'Week 1', 'prerequisites': [], 'policy_rule': item.get('policy_facts'),
                'competency': title, 'classification': 'Must Complete', 'priority': 'High'}
            db.requirements.insert_one(requirement); requirements.append(requirement)
        validation = validate_plan(FullOnboardingPlan.model_validate(content), requirements,
            list(db.documents.find()), list(db.source_sections.find()), 'support')
        assert validation['core_passed'], validation
        plan = {'_id': 'demo-plan', 'employee_id': 'learner', 'role_id': 'support', 'status': 'Review required',
            'content': content, 'validation': validation, 'learning_checks': learning_checks(content),
            'source_document_ids': [r['document_id'] for r in requirements], 'matrix_snapshot': requirements,
            'snapshot_digest': fingerprint(requirements), 'employee_snapshot': {k: employee[k] for k in ('role_id','department','experience','joining_date')},
            'created_at': now(), 'created_by': 'admin', 'generation': evidence['initial']['generation'],
            'model': evidence['model'], 'prompt_version': 'generate_v3', 'origin': 'ai_generation',
            'replay_note': 'Saved AI output; fictional local recording, not a fresh generation.'}
        db.plans.insert_one(plan)
        audit(db, 'admin', 'demo.fixture_loaded', 'demo-plan', {'note': plan['replay_note']})
        print('Isolated fictional demo ready: http://127.0.0.1:8001/app', flush=True)
        uvicorn.run(create_app(settings), host='127.0.0.1', port=8001)
    finally:
        assert db.name == schema and db.name.startswith('skillsprint_video_')
        client.drop_database(schema)
        client.close()


if __name__ == '__main__':
    main()
